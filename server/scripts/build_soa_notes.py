"""Build the draft SOA Notes (.docx): the adviser's brief to the paraplanner plus
the Statement of Advice notes, laid out like the firm's own.

    python server/scripts/build_soa_notes.py <notes.json> <out.docx>

JSON:
  brief         "Hey Howard,\n\n..."   the covering instructions
  to_obtain     ["Quotes: ...", ...]   documents the paraplanner must get (checkboxes)
  outstanding   ["Jane's DOB", ...]      facts still TBC from the meeting (checkboxes)
  circumstances "Incomes -\n..."       the short financial summary under its heading
  title         "Statement of Advice Notes - Jane and John"
  sections      [{"heading": "Goals", "body": "- ..."}, ...]

Body text: "\n" splits paragraphs, "- " is a bullet, "  - " a sub-bullet.
TBC and [ADVISER TO CONFIRM] are highlighted yellow.
"""
import json
import re
import sys
from pathlib import Path

import docx
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Pt, RGBColor

FLAG = re.compile(r'(TBC|\[ADVISER TO CONFIRM[^\]]*\])')
NAVY = RGBColor(0x1F, 0x3A, 0x5F)


def add_runs(p, text, bold=False):
    for part in filter(None, FLAG.split(text)):
        r = p.add_run(part)
        r.bold = bold
        if FLAG.fullmatch(part):
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW


def add_body(doc, text):
    for line in str(text or '').split('\n'):
        if line.startswith('  - '):
            p = doc.add_paragraph(style='List Bullet 2')
            add_runs(p, line[4:])
        elif line.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            add_runs(p, line[2:])
        else:
            p = doc.add_paragraph()
            add_runs(p, line)


def add_checklist(doc, heading, items):
    if not items:
        return
    h = doc.add_paragraph()
    add_runs(h, heading, bold=True)
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Pt(14)
        add_runs(p, '☐  ' + item)


def heading(doc, text, size=13):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(size)
    r.font.color.rgb = NAVY


def build(d, out):
    doc = docx.Document()
    st = doc.styles['Normal']
    st.font.name = 'Calibri'
    st.font.size = Pt(11)
    st.paragraph_format.space_after = Pt(4)

    add_body(doc, d.get('brief', ''))
    add_checklist(doc, 'Documents to obtain', d.get('to_obtain', []))
    add_checklist(doc, 'Still to confirm with the client', d.get('outstanding', []))
    if d.get('circumstances'):
        heading(doc, 'Financial circumstances', 12)
        add_body(doc, d['circumstances'])
    heading(doc, d.get('title', 'Statement of Advice Notes'), 15)
    for s in d.get('sections', []):
        heading(doc, s['heading'])
        add_body(doc, s.get('body', ''))
    doc.save(out)


if __name__ == '__main__':
    if len(sys.argv) < 3:
        print(__doc__)
        sys.exit(2)
    build(json.loads(Path(sys.argv[1]).read_text(encoding='utf-8')), sys.argv[2])
    print(f'wrote {sys.argv[2]}')

"""Turn an SOA pack's `0 - READ ME ... .md` into a matching .docx (and .pdf via Word).

    python legacy/packs/readme_to_docx.py "<path to READ ME .md>"

Handles the subset the READ MEs use: #/## headings, pipe tables, - bullets, 1. lists and **bold**.
"""
import os
import re
import subprocess
import sys

import docx
from docx.shared import Pt


def runs(p, text):
    text = text.replace('`', '')
    for i, part in enumerate(re.split(r'\*\*(.+?)\*\*', text)):
        if part:
            p.add_run(part).bold = i % 2 == 1


def convert(md_path):
    src = open(md_path, encoding='utf-8').read().splitlines()
    d = docx.Document()
    d.styles['Normal'].font.name = 'Arial'
    d.styles['Normal'].font.size = Pt(10)
    i = 0
    while i < len(src):
        line = src[i]
        if line.startswith('|'):
            rows = []
            while i < len(src) and src[i].startswith('|'):
                if not re.match(r'^\|[\s\-|]+\|$', src[i]):
                    rows.append([c.strip() for c in src[i].strip('|').split('|')])
                i += 1
            t = d.add_table(rows=len(rows), cols=len(rows[0]))
            t.style = 'Table Grid'
            for r, row in enumerate(rows):
                for c, val in enumerate(row):
                    cell = t.cell(r, c)
                    cell.text = ''
                    runs(cell.paragraphs[0], f'**{val}**' if r == 0 and val else val)
            d.add_paragraph()
            continue
        if line.startswith('# '):
            d.add_heading(line[2:], 0)
        elif line.startswith('## '):
            d.add_heading(line[3:], 1)
        elif line.startswith('- '):
            runs(d.add_paragraph(style='List Bullet'), line[2:])
        elif re.match(r'^\d+\. ', line):
            runs(d.add_paragraph(style='List Number'), re.sub(r'^\d+\. ', '', line))
        elif line.strip():
            runs(d.add_paragraph(), line)
        i += 1
    out = re.sub(r'\.md$', '.docx', md_path)
    d.save(out)
    return out


def to_pdf(docx_path):
    docx_path = os.path.normpath(os.path.abspath(docx_path))  # Word wants backslashes
    # narrow the "#" column of the contents table, then export
    ps = (
        f'$f="{docx_path}"; $w=New-Object -ComObject Word.Application; $w.Visible=$false; '
        'try { $d=$w.Documents.Open($f); foreach ($t in $d.Tables) { '
        'if ($t.Cell(1,1).Range.Text -like "#*" -and $t.Columns.Count -eq 3) { '
        '$t.Columns.Item(1).Width=28; $t.Columns.Item(2).Width=200; $t.Columns.Item(3).Width=220 } }; '
        '$d.Save(); $d.ExportAsFixedFormat(($f -replace "\\.docx$",".pdf"),17); $d.Close() } finally { $w.Quit() }'
    )
    subprocess.run(['powershell.exe', '-NoProfile', '-Command', ps], check=True)


if __name__ == '__main__':
    out = convert(sys.argv[1])
    to_pdf(out)
    print('wrote', out, 'and .pdf')
